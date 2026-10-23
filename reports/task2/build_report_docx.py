"""Ghép bản thảo reports/task2/01–04_*.md thành báo cáo Word Task 2 (tiêu đề, đoạn, gạch đầu dòng, bảng, khối code, hình).

Chạy từ thư mục repo: python reports/task2/build_report_docx.py [out.docx]
"""
import re, sys, os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

OUT = sys.argv[1] if len(sys.argv) > 1 else "reports/task2/BaoCao_Task2_DeTai1.docx"
PARTS = ["01_du_lieu.md", "02_baseline_asr.md", "03_thuc_nghiem.md", "04_he_thong_kha_thi.md"]
FIGS = {"01_du_lieu.md": [("data/figures/fleurs_gender.png", "Hình: nhãn giới tính FLEURS theo tập"),
                          ("data/figures/fleurs_sentence_balance.png", "Hình: số bản thu theo nhóm câu, trước và sau cân bằng")],
        "02_baseline_asr.md": [("asr/tone_confusion_FLEURS.png", "Hình: ma trận nhầm dấu thanh của PhoWhisper-small trên FLEURS test")],
        "03_thuc_nghiem.md": [("figures/curves_V1.png", "Hình: loss và BLEU dev theo epoch của V1"),
                              ("figures/ablation_bleu.png", "Hình: BLEU qua từng phiên bản")],
        "04_he_thong_kha_thi.md": [("figures/architecture.png", "Hình: kiến trúc hệ thống")]}
NAVY = RGBColor(0x0B, 0x2A, 0x6F)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri"); st.font.size = Pt(11)
for lvl, size in ((1, 16), (2, 13), (3, 12)):
    h = doc.styles[f"Heading {lvl}"]; h.font.name = "Calibri"; h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h.font.size = Pt(size); h.font.color.rgb = NAVY


def runs(p, text):
    """**đậm**, *nghiêng*, `code` trong một dòng."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10)
        elif part.startswith("*") and len(part) > 2:
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


def shade(cell, hex_):
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_)
    cell._tc.get_or_add_tcPr().append(s)


def table(lines):
    rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines if not re.match(r"^\s*\|?\s*-{3}", l)]
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        for j, v in enumerate(r[: len(rows[0])]):
            c = t.cell(i, j); c.text = ""
            p = c.paragraphs[0]; runs(p, v)
            for run in p.runs:
                run.font.size = Pt(9.5)
                if i == 0:
                    run.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if i == 0:
                shade(c, "0B2A6F")
    doc.add_paragraph()


title = doc.add_paragraph(); r = title.add_run("BÁO CÁO TASK 2 — ĐỀ TÀI 1"); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY
doc.add_paragraph("Hệ thống nhận dạng giọng nói (ASR) tiếng Việt kèm dịch sang tiếng Anh · DPL302m")
doc.add_paragraph("Xử lý dữ liệu · Thực nghiệm 2 Baseline · Proposed Method V1 → V3 · Đánh giá khả thi").runs[0].italic = True
doc.add_paragraph("Thành viên: Dương Gia Bảo (SE205262) · Lê Tuấn Kiệt (SE205279) · Hồ Lâm Bảo Đăng (SE205148) · Bùi Ngọc Phương Uyên (SE205253)")
doc.add_paragraph("Code, số liệu, log và mã hash: github.com/lekithaha06/DPL302m (thư mục reports/task2). Mọi con số lấy từ file kết quả chạy thật.")

for part in PARTS:
    lines = open(f"reports/task2/{part}", encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        l = lines[i]
        if l.startswith("```"):
            j = i + 1; block = []
            while j < len(lines) and not lines[j].startswith("```"):
                block.append(lines[j]); j += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(block)); r.font.name = "Consolas"; r.font.size = Pt(9.5)
            i = j + 1; continue
        if l.startswith("|"):
            j = i
            while j < len(lines) and lines[j].startswith("|"):
                j += 1
            table(lines[i:j]); i = j; continue
        m = re.match(r"^(#{1,3})\s+(.*)", l)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif re.match(r"^\s*[-*]\s+", l):
            p = doc.add_paragraph(style="List Bullet" if not l.startswith("  ") else "List Bullet 2")
            runs(p, re.sub(r"^\s*[-*]\s+", "", l))
        elif re.match(r"^\d+\.\s+", l):
            p = doc.add_paragraph(style="List Number"); runs(p, re.sub(r"^\d+\.\s+", "", l))
        elif l.strip():
            runs(doc.add_paragraph(), l)
        i += 1
    for path, cap in FIGS.get(part, []):
        full = f"reports/task2/{path}"
        if os.path.exists(full):
            doc.add_picture(full, width=Cm(15.5))
            c = doc.add_paragraph(); r = c.add_run(cap); r.italic = True; r.font.size = Pt(9.5)

doc.add_heading("Ghi nhận", 1)
doc.add_paragraph("Có dùng công cụ AI hỗ trợ viết script, chạy thí nghiệm và soạn bản thảo; thành viên phụ trách từng phần đọc lại, "
                  "kiểm tra số liệu với file kết quả và phải giải thích được phần đó khi vấn đáp.")
doc.save(OUT)
print("saved", OUT)
